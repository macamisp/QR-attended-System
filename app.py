import os
import io
import qrcode
from flask import Flask, render_template, request, send_file, redirect, url_for
from datetime import datetime
from openpyxl import Workbook, load_workbook
from filelock import FileLock

app = Flask(__name__)

# --- Configuration ---
# Get the absolute path of the current script
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
# Define the directory to store attendance sheets
ATTENDANCE_DIR = os.path.join(BASE_DIR, 'attendance_sheets')
# Ensure the attendance directory exists
os.makedirs(ATTENDANCE_DIR, exist_ok=True)


# --- Helper Functions ---
def get_or_create_workbook(batch_number):
    """
    Retrieves or creates an Excel workbook for the given batch number.

    Args:
        batch_number (str): The batch number for which to get or create the workbook.

    Returns:
        tuple: A tuple containing the workbook object and the filepath.
    """
    # Sanitize the batch number to create a valid filename
    sanitized_batch = "".join(c for c in batch_number if c.isalnum() or c in ('-', '_')).rstrip()
    filepath = os.path.join(ATTENDANCE_DIR, f'attendance_{sanitized_batch}.xlsx')

    if os.path.exists(filepath):
        workbook = load_workbook(filepath)
    else:
        workbook = Workbook()
        # Create a new sheet with a specific name or use the default one
        sheet = workbook.active
        sheet.title = "Attendance"
        # Add headers to the new sheet
        headers = ["Student ID", "Date", "Time", "Status"]
        sheet.append(headers)
        workbook.save(filepath)

    return workbook, filepath


# --- Routes ---
@app.route('/', methods=['GET', 'POST'])
def index():
    """
    Handles the QR code generation.
    - GET: Displays the form to enter the batch number.
    - POST: Generates a QR code for the given batch number and displays it.
    """
    if request.method == 'POST':
        batch_number = request.form['batch_number']
        # Generate the URL for the attendance page
        attendance_url = request.url_root + 'attendance/' + batch_number
        # Create a QR code instance
        qr = qrcode.QRCode(
            version=1,
            error_correction=qrcode.constants.ERROR_CORRECT_L,
            box_size=10,
            border=4,
        )
        qr.add_data(attendance_url)
        qr.make(fit=True)
        # Create an image from the QR Code instance
        img = qr.make_image(fill_color="black", back_color="white")

        # Save the image to a in-memory buffer
        buf = io.BytesIO()
        img.save(buf)
        buf.seek(0)

        # Display the QR code image
        return send_file(buf, mimetype='image/png')

    # Render the index page with the form
    return render_template('index.html')


@app.route('/attendance/<batch_number>', methods=['GET', 'POST'])
def attendance(batch_number):
    """
    Handles the student attendance marking.
    - GET: Displays the form for students to enter their ID.
    - POST: Records the student's attendance in the corresponding Excel sheet.
    """
    if request.method == 'POST':
        student_id = request.form['student_id']
        # Get the current time
        now = datetime.now()
        current_time = now.strftime("%H:%M:%S")
        current_date = now.strftime("%Y-%m-%d")

        # Define the lock file path
        lock_path = os.path.join(ATTENDANCE_DIR, f'attendance_{batch_number}.xlsx.lock')
        lock = FileLock(lock_path, timeout=10)

        with lock:
            # Get or create the workbook for the batch
            workbook, filepath = get_or_create_workbook(batch_number)
            sheet = workbook.active

            # Check for previous entry on the same day
            last_status = None
            for row in reversed(list(sheet.iter_rows(min_row=2, values_only=True))):
                if row[0] == student_id and row[1] == current_date:
                    last_status = row[3]
                    break

            status = "OUT" if last_status == "IN" else "IN"

            # Add the new attendance record
            sheet.append([student_id, current_date, current_time, status])
            workbook.save(filepath)

        # Redirect to a success page or display a message
        return "Attendance marked successfully!"

    # Render the attendance page
    return render_template('attendance.html', batch_number=batch_number)
